"""Smart Inventory - Streamlit manager interface."""

from datetime import date
from io import BytesIO
import pandas as pd
import plotly.express as px
import streamlit as st
from docx import Document

import analytics
import calendar_context as cc
import database as db
import footy


# ---------------------------------------------------------------------------
# App setup
# ---------------------------------------------------------------------------

st.set_page_config(
    page_title="Smart Inventory",
    page_icon="📦",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
      .block-container {max-width: 1450px; padding-top: 1.25rem; padding-bottom: 3rem;}
      [data-testid="stSidebar"] {background: #f7f8fb; border-right: 1px solid #e7e9ef;}
      [data-testid="stSidebar"] .block-container {padding-top: 1.2rem;}
      h1, h2, h3 {letter-spacing: -0.02em;}
      .hero {padding: 1.25rem 1.35rem; border: 1px solid #e7e9ef; border-radius: 16px; background: white; margin-bottom: 1rem;}
      .status-good {padding: 1.2rem 1.35rem; border: 1px solid #b8dfc4; border-radius: 16px; background: #f2fbf5;}
      .status-warn {padding: 1.2rem 1.35rem; border: 1px solid #f3c9a8; border-radius: 16px; background: #fff8f1;}
      .status-bad {padding: 1.2rem 1.35rem; border: 1px solid #f0b4b4; border-radius: 16px; background: #fff5f5;}
      .small-muted {color:#6b7280; font-size:0.9rem;}
      div[data-testid="stMetric"] {background:white; border:1px solid #e7e9ef; padding:1rem; border-radius:14px;}
      .stDataFrame {border:1px solid #e7e9ef; border-radius:12px; overflow:hidden;}
      .section-note {padding: .8rem 1rem; border-left: 4px solid #d7dbe5; background:#fafbfc; border-radius:8px;}

      /* Sidebar navigation as clean menu buttons */
      [data-testid="stSidebar"] div[role="radiogroup"] > label {
          padding: 0.65rem 0.8rem;
          margin-bottom: 0.25rem;
          border-radius: 10px;
          cursor: pointer;
      }
      [data-testid="stSidebar"] div[role="radiogroup"] > label:hover {
          background: #eef1f5;
      }
      [data-testid="stSidebar"] div[role="radiogroup"] > label:has(input:checked) {
          background: white;
          border: 1px solid #dfe3e8;
          font-weight: 700;
      }
      [data-testid="stSidebar"] div[role="radiogroup"] input {
          display: none;
      }
    </style>
    """,
    unsafe_allow_html=True,
)

try:
    db.init_db()
except Exception as exc:
    st.error("Could not connect to PostgreSQL.")
    st.caption("Check that PostgreSQL is running and that your .env file contains the correct database settings.")
    st.code(str(exc))
    st.stop()


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def format_qty(value, unit):
    value = float(value or 0)
    if abs(value - round(value)) < 0.05:
        value = int(round(value))
    else:
        value = round(value, 1)
    return f"{value} {unit or 'units'}"


@st.cache_data(ttl=60, show_spinner=False)
def cached_forecast(location):
    return analytics.predict_reorder_quantities(location=location)


@st.cache_data(ttl=120, show_spinner=False)
def cached_events(location):
    return analytics.get_upcoming_events(location=location, days_ahead=7)


def clear_forecast_cache():
    analytics.clear_cache()
    cached_forecast.clear()
    cached_events.clear()


def product_status_label(row):
    status = row.get("usual_order_status")
    if status == "shortfall":
        return "⚠ Needs extra"
    if status == "overorder":
        return "⚠ Reduce order"
    if status == "not_configured":
        return "⚙ Set usual order"
    return "✓ Usual order OK"


def build_order_docx(order):
    doc = Document()
    doc.add_heading(f"Purchase Order - {order['supplier_name']}", level=1)
    doc.add_paragraph(f"Order status: {order['status'].title()}")
    doc.add_paragraph(f"Created: {order['created_at']}")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Product"
    hdr[1].text = "Quantity"
    hdr[2].text = "Unit"
    for item in order["items"]:
        cells = table.add_row().cells
        cells[0].text = str(item["product_name"])
        cells[1].text = str(item["quantity"])
        cells[2].text = str(item.get("unit") or "units")
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def build_delivery_check_docx(order):
    """Create a simple receiving checklist for a placed supplier order."""
    doc = Document()
    doc.add_heading("Delivery Check Slip", level=1)
    doc.add_paragraph(f"Supplier: {order['supplier_name']}")
    doc.add_paragraph(f"Order number: {order['order_id']}")
    placed_at = order.get("sent_at") or order.get("created_at") or ""
    doc.add_paragraph(f"Order placed: {placed_at}")
    doc.add_paragraph("Delivery date: ____________________")
    doc.add_paragraph("Checked by: ______________________")
    doc.add_paragraph(
        "Use this slip when the delivery arrives. Compare each item with the order before marking the order as received in the system."
    )

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Product"
    hdr[1].text = "Ordered Qty"
    hdr[2].text = "Unit"
    hdr[3].text = "Received Qty"
    hdr[4].text = "Missing / Damaged"
    hdr[5].text = "Checked"

    for item in order["items"]:
        cells = table.add_row().cells
        cells[0].text = str(item["product_name"])
        cells[1].text = str(item["quantity"])
        cells[2].text = str(item.get("unit") or "units")
        cells[3].text = "________"
        cells[4].text = "________"
        cells[5].text = "[  ]"

    doc.add_paragraph()
    doc.add_paragraph("All items received correctly:  [  ] Yes   [  ] No")
    doc.add_paragraph("Notes: __________________________________________________________")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def build_supplier_delivery_check_docx(supplier_name, sent_orders):
    """Create one receiving checklist for all open placed orders from a supplier."""
    doc = Document()
    doc.add_heading("Delivery Check Slip", level=1)
    doc.add_paragraph(f"Supplier: {supplier_name}")

    order_numbers = ", ".join(str(o["order_id"]) for o in sent_orders)
    doc.add_paragraph(f"Order number(s): {order_numbers}")

    placed_times = [str(o.get("sent_at") or o.get("created_at") or "") for o in sent_orders]
    placed_times = [v for v in placed_times if v]
    if placed_times:
        doc.add_paragraph(f"Order placed: {', '.join(placed_times)}")

    doc.add_paragraph("Delivery date: ____________________")
    doc.add_paragraph("Checked by: ______________________")
    doc.add_paragraph(
        "Check each product against the quantity ordered before confirming the delivery in the system."
    )

    # Combine duplicate products if older data contains more than one open
    # placed order for the same supplier.
    combined = {}
    for order in sent_orders:
        for item in order.get("items", []):
            key = (int(item["product_id"]), str(item.get("unit") or "units"))
            if key not in combined:
                combined[key] = {
                    "product_name": item["product_name"],
                    "quantity": 0,
                    "unit": item.get("unit") or "units",
                }
            combined[key]["quantity"] += int(item.get("quantity") or 0)

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Product"
    hdr[1].text = "Ordered Qty"
    hdr[2].text = "Unit"
    hdr[3].text = "Received Qty"
    hdr[4].text = "Missing / Damaged"
    hdr[5].text = "Checked"

    for item in sorted(combined.values(), key=lambda x: x["product_name"].lower()):
        cells = table.add_row().cells
        cells[0].text = str(item["product_name"])
        cells[1].text = str(item["quantity"])
        cells[2].text = str(item["unit"])
        cells[3].text = "________"
        cells[4].text = "________"
        cells[5].text = "[  ]"

    doc.add_paragraph()
    doc.add_paragraph("All items received correctly:  [  ] Yes   [  ] No")
    doc.add_paragraph("Notes: __________________________________________________________")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def safe_forecast(location):
    try:
        return cached_forecast(location)
    except Exception as exc:
        st.warning("Forecast unavailable right now. Inventory and supplier data are still available.")
        st.caption(str(exc))
        return [], {}


def safe_events(location):
    try:
        return cached_events(location)
    except Exception as exc:
        st.warning("Upcoming event information could not be refreshed.")
        st.caption(str(exc))
        return [], {}


def safe_footy_matches():
    try:
        return footy.get_year_matches()
    except Exception:
        return []


# ---------------------------------------------------------------------------
# Sidebar navigation
# ---------------------------------------------------------------------------

st.sidebar.markdown("## 📦 Smart Inventory")
st.sidebar.caption("Inventory, suppliers and weekly ordering")

# This project is configured for one Adelaide store, so there is no region selector.
selected_location = "Adelaide, SA"

PAGE_LABELS = {
    "🏠  Home": "Home",
    "🚚  Suppliers": "Suppliers",
    "📦  Products": "Products",
    "📝  Daily Stock Update": "Daily Stock Update",
    "📅  Upcoming Events": "Upcoming Events",
    "🏉  Footy Match Alerts": "Footy Match Alerts",
    "🛒  Weekly Orders": "Orders",
}

selected_page = st.sidebar.radio(
    "Navigation",
    list(PAGE_LABELS.keys()),
    label_visibility="collapsed",
)
page = PAGE_LABELS[selected_page]


# ---------------------------------------------------------------------------
# HOME
# ---------------------------------------------------------------------------

if page == "Home":
    st.title("Weekly order check")
    st.caption("One decision: is your normal weekly order still safe for next week?")

    inventory = db.get_manager_inventory()
    forecast, meta = safe_forecast(selected_location)

    if not inventory:
        st.markdown(
            """
            <div class="hero">
              <h3>No products yet</h3>
              <p>Add your products and suppliers first. The dashboard will then compare your actual inventory, usual order and forecast automatically.</p>
            </div>
            """,
            unsafe_allow_html=True,
        )
        st.info("Open the Products page to add your first product.")
    else:

        attention = [r for r in forecast if r.get("usual_order_status") != "good"]

        if forecast and not attention:
            st.markdown(
                """
                <div class="status-good">
                  <h2>✓ All good</h2>
                  <p>Your current inventory plus the usual weekly order should cover next week's predicted demand without exceeding the configured maximum stock levels.</p>
                  <p><b>Proceed with the usual order.</b></p>
                </div>
                """,
                unsafe_allow_html=True,
            )
        elif forecast:
            shortages = [r for r in attention if r["usual_order_status"] in {"shortfall", "not_configured"}]
            overorders = [r for r in attention if r["usual_order_status"] == "overorder"]
            st.markdown(
                f"""
                <div class="status-bad">
                  <h2>⚠ Order needs attention</h2>
                  <p>Your usual order should be changed for <b>{len(attention)} product(s)</b> before next week.</p>
                  <p>{len(shortages)} shortage/configuration issue(s) · {len(overorders)} possible overorder issue(s)</p>
                </div>
                """,
                unsafe_allow_html=True,
            )

            rows = []
            for r in attention:
                usual = int(r["usual_order_qty"])
                recommended = int(r["recommended_final_order_qty"])
                change = recommended - usual
                if change > 0:
                    change_text = f"+{change} {r['unit']}"
                elif change < 0:
                    change_text = f"{change} {r['unit']}"
                else:
                    change_text = "Configure"
                rows.append({
                    "Product": r["name"],
                    "Supplier": r["supplier_name"],
                    "Usual order": format_qty(usual, r["unit"]),
                    "Recommended": format_qty(recommended, r["unit"]),
                    "Change": change_text,
                    "Reason": r["reason"],
                })
            st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)
            st.info("Open Orders to review the safe quantities before placing the supplier order.")
        else:
            st.info("There is not enough sales/usage history yet to produce a weekly order check.")


# ---------------------------------------------------------------------------
# SUPPLIERS
# ---------------------------------------------------------------------------

elif page == "Suppliers":
    st.title("Suppliers")
    st.caption("Select a supplier to see exactly what they supply, their costs and which products use them as the primary supplier.")

    suppliers = db.get_suppliers()

    top_left, top_right = st.columns([2.2, 1], gap="large")

    with top_left:
        if suppliers:
            supplier_map = {s["name"]: s for s in suppliers}
            selected_supplier_name = st.selectbox("Select supplier", list(supplier_map.keys()))
            supplier = supplier_map[selected_supplier_name]

            d1, d2, d3 = st.columns(3)
            d1.metric("Delivery time", f"{supplier.get('lead_time_days') or 0} days")
            d2.metric("Products supplied", len(db.get_supplier_catalog(supplier["supplier_id"])))
            d3.metric("Status", "Active")

            st.subheader("Products supplied")
            catalog = db.get_supplier_catalog(supplier["supplier_id"])
            if catalog:
                catalog_df = pd.DataFrame([
                    {
                        "Product": x["product_name"],
                        "Category": x["category"],
                        "Cost": "" if x["unit_cost"] is None else f"${float(x['unit_cost']):.2f}",
                        "Primary supplier": "Yes" if x["is_primary"] else "No",
                    }
                    for x in catalog
                ])
                st.dataframe(catalog_df, use_container_width=True, hide_index=True)
            else:
                st.info("No products are linked to this supplier yet.")

            with st.expander("Link another product to this supplier", expanded=not bool(catalog)):
                available = db.get_unlinked_products_for_supplier(supplier["supplier_id"])
                if not available:
                    st.caption("All active products are already linked to this supplier.")
                else:
                    pmap = {f"{p['name']} · {p['category']}": p for p in available}
                    label = st.selectbox("Product", list(pmap.keys()), key="supplier_link_product")
                    product = pmap[label]
                    cost = st.number_input("Cost from this supplier", min_value=0.0, value=0.0, step=0.50)
                    make_primary = st.checkbox("Make this the primary supplier")
                    if st.button("Add supplied product", type="primary"):
                        db.link_product_supplier(
                            product["product_id"],
                            supplier["supplier_id"],
                            cost,
                            is_primary=make_primary,
                        )
                        clear_forecast_cache()
                        st.success("Product linked to supplier.")
                        st.rerun()

            if catalog:
                with st.expander("Change a supplied product"):
                    link_map = {x["product_name"]: x for x in catalog}
                    product_name = st.selectbox("Supplied product", list(link_map.keys()), key="supplier_edit_link")
                    link = link_map[product_name]
                    new_cost = st.number_input(
                        "Cost",
                        min_value=0.0,
                        value=float(link["unit_cost"] or 0),
                        step=0.50,
                        key="supplier_link_cost",
                    )
                    c1, c2, c3 = st.columns(3)
                    if c1.button("Save cost"):
                        db.link_product_supplier(
                            link["product_id"], supplier["supplier_id"], new_cost,
                            is_primary=bool(link["is_primary"]),
                        )
                        st.rerun()
                    if c2.button("Make primary", disabled=bool(link["is_primary"])):
                        db.set_primary_supplier(link["product_id"], supplier["supplier_id"])
                        clear_forecast_cache()
                        st.rerun()
                    if c3.button("Remove link"):
                        db.unlink_product_supplier(link["product_id"], supplier["supplier_id"])
                        clear_forecast_cache()
                        st.rerun()
        else:
            st.info("No suppliers yet. Add your first supplier using the form on the right.")

    with top_right:
        st.subheader("Manage supplier")
        with st.expander("+ Add supplier", expanded=not bool(suppliers)):
            with st.form("add_supplier_form", clear_on_submit=True):
                name = st.text_input("Supplier name")
                email = st.text_input("Email")
                phone = st.text_input("Phone")
                lead = st.number_input("Delivery time (days)", min_value=0, value=3, step=1)
                if st.form_submit_button("Add supplier", type="primary"):
                    db.add_supplier(name, email, phone, int(lead))
                    clear_forecast_cache()
                    st.success("Supplier added to PostgreSQL.")
                    st.rerun()

        if suppliers:
            with st.expander("Edit selected supplier"):
                with st.form("edit_supplier_form"):
                    edit_name = st.text_input("Name", value=supplier["name"])
                    edit_email = st.text_input("Email", value=supplier.get("contact_email") or "")
                    edit_phone = st.text_input("Phone", value=supplier.get("phone") or "")
                    edit_lead = st.number_input(
                        "Delivery time (days)", min_value=0,
                        value=int(supplier.get("lead_time_days") or 0), step=1,
                    )
                    if st.form_submit_button("Save supplier"):
                        db.update_supplier(supplier["supplier_id"], edit_name, edit_email, edit_phone, edit_lead)
                        clear_forecast_cache()
                        st.rerun()

            with st.expander("Remove selected supplier"):
                linked = db.get_supplier_primary_products(supplier["supplier_id"])
                if linked:
                    st.warning("Move these primary products to another supplier first: " + ", ".join(p["name"] for p in linked))
                else:
                    remove_mode = st.radio("Remove type", ["Archive", "Delete permanently"], horizontal=True)
                    if st.button("Remove supplier"):
                        if remove_mode == "Archive":
                            db.set_supplier_active(supplier["supplier_id"], False)
                        else:
                            db.delete_supplier_permanently(supplier["supplier_id"])
                        clear_forecast_cache()
                        st.rerun()


# ---------------------------------------------------------------------------
# PRODUCTS
# ---------------------------------------------------------------------------

elif page == "Products":
    st.title("Products")
    st.caption("Browse products by category and manage the ordering rules that drive the weekly forecast.")

    products = db.get_products_admin(include_inactive=False)
    suppliers = db.get_suppliers()
    supplier_map = {s["name"]: s for s in suppliers}
    categories = sorted({p.get("category") or "Other" for p in products})

    browse_col, manage_col = st.columns([2.2, 1], gap="large")

    with browse_col:
        if products:
            category = st.selectbox("Category", ["All categories"] + categories)
            visible = products if category == "All categories" else [p for p in products if (p.get("category") or "Other") == category]
            product_df = pd.DataFrame([
                {
                    "Product": p["name"],
                    "Category": p.get("category") or "Other",
                    "Unit": p.get("unit") or "units",
                    "Current stock": p["current_stock"],
                    "Safety stock": p["reorder_threshold"],
                    "Usual weekly order": p["usual_order_qty"],
                    "Maximum stock": p["max_stock_level"] if p["max_stock_level"] is not None else "Not set",
                    "Primary supplier": p.get("primary_supplier_name") or "No supplier",
                }
                for p in visible
            ])
            st.dataframe(product_df, use_container_width=True, hide_index=True)

            st.subheader("Edit product")
            pmap = {p["name"]: p for p in products}
            selected_name = st.selectbox("Choose product", list(pmap.keys()))
            product = pmap[selected_name]
            current_supplier = product.get("primary_supplier_name") or "No supplier"
            supplier_options = ["No supplier"] + list(supplier_map.keys())
            if current_supplier not in supplier_options:
                supplier_options.append(current_supplier)

            with st.form("edit_product_form"):
                c1, c2 = st.columns(2)
                name = c1.text_input("Product name", value=product["name"])
                category_value = c2.text_input("Category", value=product.get("category") or "Other")
                unit = c1.text_input("Unit", value=product.get("unit") or "units")
                safety = c2.number_input("Safety stock", min_value=0, value=int(product["reorder_threshold"]), step=1)
                usual = c1.number_input("Usual weekly order", min_value=0, value=int(product["usual_order_qty"] or 0), step=1)
                max_stock = c2.number_input("Maximum stock (0 = not set)", min_value=0, value=int(product["max_stock_level"] or 0), step=1)
                shelf = c1.number_input("Shelf life days (0 = not set)", min_value=0, value=int(product["shelf_life_days"] or 0), step=1)
                supplier_name = c2.selectbox("Primary supplier", supplier_options, index=supplier_options.index(current_supplier))
                unit_cost = c1.number_input("Cost from primary supplier", min_value=0.0, value=float(product.get("unit_cost") or 0), step=0.50)
                if st.form_submit_button("Save product", type="primary"):
                    supplier_id = None if supplier_name == "No supplier" else supplier_map[supplier_name]["supplier_id"]
                    db.update_manager_product(
                        product["product_id"], name, category_value, unit, safety, usual,
                        max_stock_level=max_stock, shelf_life_days=shelf,
                        supplier_id=supplier_id, unit_cost=unit_cost,
                    )
                    clear_forecast_cache()
                    st.success("Product updated in PostgreSQL.")
                    st.rerun()
        else:
            st.info("No products yet. Add your first product using the form on the right.")

    with manage_col:
        st.subheader("Add / remove")
        with st.expander("+ Add product", expanded=not bool(products)):
            with st.form("add_product_form", clear_on_submit=True):
                name = st.text_input("Product name")
                category = st.text_input("Category")
                unit = st.text_input("Unit", value="units")
                stock = st.number_input("Current stock", min_value=0.0, value=0.0, step=0.5)
                safety = st.number_input("Safety stock", min_value=0, value=0, step=1)
                usual = st.number_input("Usual weekly order", min_value=0, value=0, step=1)
                max_stock = st.number_input("Maximum stock (0 = not set)", min_value=0, value=0, step=1)
                shelf = st.number_input("Shelf life days (0 = not set)", min_value=0, value=0, step=1)
                supplier_name = st.selectbox("Primary supplier", ["No supplier"] + list(supplier_map.keys()))
                cost = st.number_input("Supplier cost", min_value=0.0, value=0.0, step=0.50)
                if st.form_submit_button("Add product", type="primary"):
                    supplier_id = None if supplier_name == "No supplier" else supplier_map[supplier_name]["supplier_id"]
                    db.add_manager_product(
                        name, category, unit, safety, stock,
                        usual_order_qty=usual,
                        max_stock_level=max_stock,
                        shelf_life_days=shelf,
                        supplier_id=supplier_id,
                        unit_cost=cost,
                    )
                    clear_forecast_cache()
                    st.success("Product added to PostgreSQL.")
                    st.rerun()

        if products:
            with st.expander("Remove product"):
                remove_name = st.selectbox("Product to remove", [p["name"] for p in products], key="remove_product_select")
                remove_product = next(p for p in products if p["name"] == remove_name)
                mode = st.radio("Remove type", ["Archive", "Delete permanently"], horizontal=True, key="remove_product_mode")
                st.caption("Permanent delete is allowed only when the product has no stock, order or nightly-count history.")
                if st.button("Remove product"):
                    if mode == "Archive":
                        db.set_product_active(remove_product["product_id"], False)
                    else:
                        db.delete_product_permanently(remove_product["product_id"])
                    clear_forecast_cache()
                    st.rerun()


# ---------------------------------------------------------------------------
# INVENTORY
# ---------------------------------------------------------------------------

elif page == "Daily Stock Update":
    st.title("Daily Stock Update")
    st.caption("Enter today's usage for every item, then confirm each row before saving.")

    usage_rows = db.get_daily_usage(date.today())

    if not usage_rows:
        st.info("No products have been added yet.")
    else:
        rows = [
            {
                "product_id": int(item["product_id"]),
                "Item": item["product_name"],
                "Used today": float(item.get("quantity_used", 0) or 0),
                "Unit": item.get("unit") or "units",
                "Updated": False,
            }
            for item in usage_rows
        ]

        edited = st.data_editor(
            pd.DataFrame(rows),
            use_container_width=True,
            hide_index=True,
            disabled=["product_id", "Item", "Unit"],
            column_config={
                "product_id": None,
                "Item": st.column_config.TextColumn("Item", width="large"),
                "Used today": st.column_config.NumberColumn(
                    "Count used today",
                    min_value=0.0,
                    step=0.5,
                    format="%.1f",
                    width="small",
                ),
                "Unit": st.column_config.TextColumn("Unit", width="small"),
                "Updated": st.column_config.CheckboxColumn(
                    "Updated",
                    help="Tick this after checking the quantity for this item.",
                    width="small",
                ),
            },
            key="daily_usage_editor",
        )

        all_updated = bool(edited["Updated"].all())
        if not all_updated:
            st.info("Check every item first. The daily update can only be saved when all rows are marked Updated.")

        if st.button(
            "Save Daily Update",
            type="primary",
            use_container_width=True,
            disabled=not all_updated,
        ):
            quantities = {
                int(row["product_id"]): float(row["Used today"] or 0)
                for _, row in edited.iterrows()
            }

            try:
                db.save_daily_usage(date.today(), quantities)
                clear_forecast_cache()
                st.success("Today's full stock update has been saved and inventory has been updated.")
                st.rerun()
            except ValueError as exc:
                st.error(str(exc))


# ---------------------------------------------------------------------------
# UPCOMING EVENTS
# ---------------------------------------------------------------------------

elif page == "Upcoming Events":
    st.title("Upcoming events & forecast reasons")
    st.caption("See the next 7 days of weekends, holidays and weather, and exactly how they affect the forecast.")

    events, event_meta = safe_events(selected_location)
    if events:
        event_df = pd.DataFrame([
            {
                "Date": e["date"].strftime("%d %b"),
                "Day": e["day"],
                "Event / condition": e["conditions"],
                "Rain": "—" if e["rain_mm"] is None else f"{e['rain_mm']} mm",
                "Max temp": "—" if e["max_temp_c"] is None else f"{e['max_temp_c']}°C",
                "Factor used": e["factor_used"],
                "Forecast effect": e["impact"],
            }
            for e in events
        ])
        st.dataframe(event_df, use_container_width=True, hide_index=True)

    forecast, meta = safe_forecast(selected_location)
    if forecast:
        st.divider()
        st.subheader("Why is a product being predicted this way?")
        fmap = {r["name"]: r for r in forecast}
        product_name = st.selectbox("Product", list(fmap.keys()))
        row = fmap[product_name]

        d1, d2, d3, d4 = st.columns(4)
        d1.metric("Have now", format_qty(row["current_stock"], row["unit"]))
        d2.metric("Predicted 7-day use", format_qty(row["predicted_week_demand"], row["unit"]))
        d3.metric("Usual order", format_qty(row["usual_order_qty"], row["unit"]))
        d4.metric("Recommended", format_qty(row["recommended_final_order_qty"], row["unit"]))

        chart_df = pd.DataFrame(row["daily_forecast"])
        fig = px.bar(
            chart_df,
            x="day",
            y="predicted_demand",
            hover_data=["factor_used", "multiplier", "rain_mm", "max_temp_c"],
            labels={"predicted_demand": f"Predicted {row['unit']}", "day": "Day"},
            title=f"{row['name']} - predicted daily usage",
        )
        st.plotly_chart(fig, use_container_width=True)

        st.markdown("#### Calculation")
        st.write(
            f"Baseline average: **{row['baseline_daily_avg']} {row['unit']} per observed sales day**. "
            "Public holidays have first priority. Otherwise the strongest applicable weekend, school-holiday or weather effect is used, with the monthly pattern used only when no special condition applies."
        )
        st.code(
            f"Forecast need = predicted 7-day use + safety stock - current stock - already on order\n"
            f"              = {row['predicted_week_demand']} + {row['safety_stock']} - {row['current_stock']} - {row['on_order_qty']}\n"
            f"              = {row['required_order_qty']} {row['unit']}"
        )

        if row.get("max_stock_level") is not None:
            st.code(
                f"Maximum stock cap = max(0, maximum stock - current stock - already on order)\n"
                f"                  = max(0, {row['max_stock_level']} - {row['current_stock']} - {row['on_order_qty']})\n"
                f"                  = {row.get('max_safe_order_qty', 0)} {row['unit']}\n\n"
                f"Recommended order = {row['recommended_final_order_qty']} {row['unit']}"
            )

        st.info(row["reason"])
        confidence = meta.get('confidence', 'Unknown')
        days_of_data = meta.get('days_of_data', 0)
        if confidence in {"Very low", "Low"}:
            st.warning(
                f"Forecast confidence is {confidence.lower()} because only {days_of_data} observed sales/usage day(s) are available. "
                "Treat the predicted demand as an early estimate until more history is collected."
            )
        else:
            st.caption(f"Forecast confidence: {confidence} · based on {days_of_data} observed sales/usage day(s).")


# ---------------------------------------------------------------------------
# FOOTY MATCH ALERTS
# ---------------------------------------------------------------------------

elif page == "Footy Match Alerts":
    year = footy.current_year()
    today = footy.current_date()
    st.title("Footy match alerts")
    st.caption(
        f"Adelaide Crows, Port Adelaide, AFL finals and SANFL matches for {year}. "
        "Times are shown in Adelaide time."
    )

    matches = safe_footy_matches()
    today_matches = [m for m in matches if m["datetime"].date() == today]
    upcoming_matches = [m for m in matches if m["datetime"].date() > today]
    past_matches = sorted(
        [m for m in matches if m["datetime"].date() < today],
        key=lambda m: m["datetime"],
        reverse=True,
    )

    def show_matches(match_list):
        rows = [
            {
                "Competition": match["competition"],
                "Round": match["round"],
                "Date": match["datetime"].strftime("%d %b %Y"),
                "Day": match["datetime"].strftime("%A"),
                "Time": match["datetime"].strftime("%I:%M %p").lstrip("0"),
                "Match": match["match"],
                "Venue": match["venue"],
            }
            for match in match_list
        ]
        st.dataframe(pd.DataFrame(rows), hide_index=True, use_container_width=True)

    with st.expander("🏉 Today's Matches", expanded=True):
        if today_matches:
            show_matches(today_matches)
        else:
            st.info("No footy matches today.")

    with st.expander("📅 Upcoming Matches", expanded=True):
        if upcoming_matches:
            show_matches(upcoming_matches)
        else:
            st.info(f"No more published upcoming matches are available for {year}.")

    with st.expander("🕘 Past Matches", expanded=True):
        if past_matches:
            show_matches(past_matches)
        else:
            st.info(f"No past matches are available for {year}.")

    st.caption(
        f"Fixtures refresh from the current-year feeds when this page is opened or reloaded. "
        f"When {year + 1} begins, the page automatically starts requesting {year + 1} fixtures."
    )


# ---------------------------------------------------------------------------
# ORDERS
# ---------------------------------------------------------------------------

elif page == "Orders":
    st.title("Weekly Orders")
    st.caption("See each supplier once, with all products and quantities that belong to that supplier.")

    forecast, _ = safe_forecast(selected_location)
    open_order_status = db.get_open_order_item_statuses()
    orders = db.get_orders_grouped()

    # Existing database orders grouped under ONE supplier section.
    orders_by_supplier = {}
    for order in orders:
        supplier_id = int(order["supplier_id"])
        group = orders_by_supplier.setdefault(
            supplier_id,
            {
                "supplier_id": supplier_id,
                "supplier_name": order["supplier_name"],
                "draft_orders": [],
                "sent_orders": [],
            },
        )
        if order["status"] == "draft":
            group["draft_orders"].append(order)
        elif order["status"] == "sent":
            group["sent_orders"].append(order)

    # Build a simple product-level plan. Products with the same supplier will
    # be displayed inside the same supplier card.
    order_plan = []
    forecast_by_product = {}
    for row in forecast or []:
        product_id = int(row["product_id"])
        forecast_by_product[product_id] = row
        recommended = int(row.get("recommended_final_order_qty", 0) or 0)
        existing = open_order_status.get(product_id)

        if existing and existing["status"] in {"draft", "sent"}:
            planned_qty = int(existing["quantity"] or 0)
        else:
            planned_qty = recommended

        if planned_qty <= 0:
            continue

        supplier_id = row.get("supplier_id")
        supplier_name = row.get("supplier_name") or "Supplier not linked"

        if not supplier_id:
            status = "Supplier needed"
        elif existing and existing["status"] == "sent":
            status = "Ordered - waiting for delivery"
        elif existing and existing["status"] == "draft":
            status = "Ready to place"
        else:
            status = "Not added yet"

        order_plan.append(
            {
                "Product": row["name"],
                "Quantity": format_qty(planned_qty, row.get("unit")),
                "Supplier": supplier_name,
                "Status": status,
                "product_id": product_id,
                "supplier_id": int(supplier_id) if supplier_id else None,
                "supplier_name": supplier_name,
            }
        )

    if not order_plan and not orders:
        st.success("There are no products to order right now.")
    else:
        view_mode = st.radio(
            "View order list by",
            ["Supplier", "Product"],
            horizontal=True,
            key="weekly_order_view_mode",
        )

        if view_mode == "Product":
            if order_plan:
                product_rows = sorted(
                    order_plan,
                    key=lambda x: (x["Product"].lower(), x["Supplier"].lower()),
                )
                st.dataframe(
                    pd.DataFrame(product_rows)[["Product", "Quantity", "Supplier", "Status"]],
                    hide_index=True,
                    use_container_width=True,
                )
                st.caption("Switch to Supplier view to prepare, place, or receive an order.")
            else:
                st.info("No products currently need ordering.")

        else:
            # Supplier view is the main working view. Every supplier appears
            # only once, even if old data contains more than one order row.
            supplier_ids = {
                int(r["supplier_id"])
                for r in order_plan
                if r.get("supplier_id")
            } | set(orders_by_supplier.keys())

            supplier_name_lookup = {
                int(r["supplier_id"]): r["supplier_name"]
                for r in order_plan
                if r.get("supplier_id")
            }
            for sid, group in orders_by_supplier.items():
                supplier_name_lookup[sid] = group["supplier_name"]

            for supplier_id in sorted(
                supplier_ids,
                key=lambda sid: supplier_name_lookup.get(sid, "").lower(),
            ):
                supplier_name = supplier_name_lookup.get(supplier_id, "Supplier")
                supplier_rows = [
                    r for r in order_plan if r.get("supplier_id") == supplier_id
                ]
                supplier_group = orders_by_supplier.get(
                    supplier_id,
                    {
                        "supplier_id": supplier_id,
                        "supplier_name": supplier_name,
                        "draft_orders": [],
                        "sent_orders": [],
                    },
                )

                with st.container(border=True):
                    st.markdown(f"### {supplier_name}")

                    # One table per supplier. It contains every product for
                    # that supplier, regardless of whether it is new, draft,
                    # or already placed.
                    if supplier_rows:
                        st.dataframe(
                            pd.DataFrame(supplier_rows)[["Product", "Quantity", "Status"]],
                            hide_index=True,
                            use_container_width=True,
                        )
                    else:
                        # Fallback for an open order not present in the current
                        # forecast result.
                        fallback_rows = []
                        for status_key, label in (
                            ("draft_orders", "Ready to place"),
                            ("sent_orders", "Ordered - waiting for delivery"),
                        ):
                            for order in supplier_group[status_key]:
                                for item in order.get("items", []):
                                    fallback_rows.append(
                                        {
                                            "Product": item["product_name"],
                                            "Quantity": format_qty(item["quantity"], item.get("unit")),
                                            "Status": label,
                                        }
                                    )
                        if fallback_rows:
                            st.dataframe(
                                pd.DataFrame(fallback_rows),
                                hide_index=True,
                                use_container_width=True,
                            )

                    new_items = [
                        r for r in supplier_rows if r["Status"] == "Not added yet"
                    ]
                    draft_orders = supplier_group["draft_orders"]
                    sent_orders = supplier_group["sent_orders"]

                    if new_items:
                        st.caption(
                            "These products belong to this supplier. Add them to the supplier's current draft order."
                        )
                        if st.button(
                            f"Add products to {supplier_name} order",
                            type="primary" if not draft_orders else "secondary",
                            key=f"supplier_add_{supplier_id}",
                            use_container_width=True,
                        ):
                            blocked = []
                            for plan_row in new_items:
                                source = forecast_by_product.get(int(plan_row["product_id"]))
                                if not source:
                                    continue
                                try:
                                    db.upsert_draft_order_item(
                                        source["product_id"],
                                        int(source.get("recommended_final_order_qty", 0) or 0),
                                    )
                                except ValueError as exc:
                                    blocked.append(f"{source['name']}: {exc}")
                            clear_forecast_cache()
                            if blocked:
                                st.warning("\n\n".join(blocked))
                            st.rerun()

                    if draft_orders:
                        draft_item_count = sum(len(o.get("items", [])) for o in draft_orders)
                        st.caption(
                            f"{draft_item_count} prepared product(s) are ready to place with {supplier_name}."
                        )
                        if st.button(
                            f"Place {supplier_name} order",
                            type="primary",
                            key=f"supplier_place_{supplier_id}",
                            use_container_width=True,
                        ):
                            for order in draft_orders:
                                db.mark_order_sent(order["order_id"])
                            clear_forecast_cache()
                            st.rerun()

                        with st.expander("Edit prepared products only if needed"):
                            draft_items = []
                            for order in draft_orders:
                                for item in order.get("items", []):
                                    draft_items.append((order, item))

                            if draft_items:
                                labels = {
                                    f"{item['product_name']} ({format_qty(item['quantity'], item.get('unit'))})": (order, item)
                                    for order, item in draft_items
                                }
                                selected_label = st.selectbox(
                                    "Product",
                                    list(labels.keys()),
                                    key=f"supplier_edit_product_{supplier_id}",
                                )
                                selected_order, selected_item = labels[selected_label]
                                new_qty = st.number_input(
                                    "Quantity",
                                    min_value=1,
                                    value=int(selected_item["quantity"]),
                                    step=1,
                                    key=f"supplier_edit_qty_{supplier_id}_{selected_item['product_id']}",
                                )
                                e1, e2 = st.columns(2)
                                if e1.button(
                                    "Save quantity",
                                    key=f"supplier_save_{supplier_id}_{selected_item['product_id']}",
                                    use_container_width=True,
                                ):
                                    db.update_order_item_quantity(
                                        selected_order["order_id"],
                                        selected_item["product_id"],
                                        int(new_qty),
                                    )
                                    clear_forecast_cache()
                                    st.rerun()
                                if e2.button(
                                    "Remove product",
                                    key=f"supplier_remove_{supplier_id}_{selected_item['product_id']}",
                                    use_container_width=True,
                                ):
                                    db.remove_order_item(
                                        selected_order["order_id"],
                                        selected_item["product_id"],
                                    )
                                    clear_forecast_cache()
                                    st.rerun()

                    if sent_orders:
                        st.info(
                            "Order placed. When the delivery arrives, use one delivery check slip for this supplier."
                        )
                        b1, b2 = st.columns(2)
                        b1.download_button(
                            "Download delivery check slip",
                            data=build_supplier_delivery_check_docx(supplier_name, sent_orders),
                            file_name=f"delivery_check_{supplier_name.replace(' ', '_')}.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            key=f"supplier_delivery_check_{supplier_id}",
                            use_container_width=True,
                        )
                        if b2.button(
                            "All ordered items received - update stock",
                            type="primary",
                            key=f"supplier_received_{supplier_id}",
                            use_container_width=True,
                        ):
                            for order in sent_orders:
                                db.mark_order_received(order["order_id"])
                            clear_forecast_cache()
                            st.success(
                                f"{supplier_name} delivery completed and received quantities were added to inventory."
                            )
                            st.rerun()

                        st.caption(
                            "If anything is missing or damaged, do not complete the delivery yet. Keep it open until it is resolved."
                        )

            missing_supplier = [r for r in order_plan if not r.get("supplier_id")]
            if missing_supplier:
                st.warning("Some products do not have a primary supplier linked yet.")
                st.dataframe(
                    pd.DataFrame(missing_supplier)[["Product", "Quantity", "Status"]],
                    hide_index=True,
                    use_container_width=True,
                )

